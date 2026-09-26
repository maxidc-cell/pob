# -*- coding: utf-8 -*-
"""
Genera Plantilla_Reportes_POB.xlsx: la hoja Config con todas las tablas con
nombre ya cargadas (usando los mismos valores por defecto que
validacion/config.py, para que Python y Excel arranquen alineados) y una
hoja vacía por cada reporte de salida.

No depende de tener Excel instalado: usa openpyxl. Se corre una sola vez
para generar el archivo base; después el usuario lo abre en su Excel
corporativo y sigue GUIA_INSTALACION.md.
"""
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from validacion import config as cfg

TITULO_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
TITULO_FONT = Font(color="FFFFFF", bold=True)
ENCABEZADO_FONT = Font(bold=True)


def agregar_tabla(ws, nombre_tabla, encabezados, filas, fila_inicio, col_inicio=1):
    """Escribe encabezados + filas a partir de (fila_inicio, col_inicio) y las
    registra como Tabla de Excel con nombre `nombre_tabla` (lo que Power
    Query lee vía Excel.CurrentWorkbook())."""
    for j, encabezado in enumerate(encabezados):
        celda = ws.cell(row=fila_inicio, column=col_inicio + j, value=encabezado)
        celda.font = ENCABEZADO_FONT
    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila):
            ws.cell(row=fila_inicio + 1 + i, column=col_inicio + j, value=valor)

    fila_fin = fila_inicio + len(filas)
    col_fin = col_inicio + len(encabezados) - 1
    ref = (
        f"{get_column_letter(col_inicio)}{fila_inicio}:"
        f"{get_column_letter(col_fin)}{max(fila_fin, fila_inicio + 1)}"
    )
    # Si no hay filas de datos, Excel igual necesita al menos una fila para
    # el rango de la tabla; dejamos una fila vacía por debajo del encabezado
    # en ese caso (ya contemplado por max() arriba: fila_inicio+1).
    tabla = Table(displayName=nombre_tabla, ref=ref)
    tabla.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False
    )
    ws.add_table(tabla)
    return fila_fin


def hoja_config(wb):
    ws = wb.active
    ws.title = "Config"

    ws["A1"] = "Configuración de los reportes POB y Comedor"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        "Editá estas tablas para ajustar el comportamiento de las consultas, sin tocar el código M. "
        "Ver GUIA_INSTALACION.md para el detalle de cada parámetro."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:F2")

    fila = 4

    ws.cell(row=fila, column=1, value="tParametros").font = Font(bold=True, italic=True)
    fila += 1
    filas_parametros = [
        ("URL_Sitio_OneDrive", ""),
        ("Ruta_Carpeta", ""),
        ("Prefijo_Archivo", cfg.PREFIJO_ARCHIVO),
        ("Hoja_Origen", cfg.HOJA_ORIGEN),
    ]
    fila = agregar_tabla(ws, "tParametros", ["Parametro", "Valor"], filas_parametros, fila) + 3

    ws.cell(row=fila, column=1, value="tHorasFoto").font = Font(bold=True, italic=True)
    fila += 1
    filas_horas = [(h,) for h in cfg.HORAS_FOTO]
    fila = agregar_tabla(ws, "tHorasFoto", ["Hora"], filas_horas, fila) + 3

    ws.cell(row=fila, column=1, value="tFranjas").font = Font(bold=True, italic=True)
    fila += 1
    filas_franjas = [(f["Franja"], f["HoraInicio"], f["HoraCierre"]) for f in cfg.FRANJAS]
    fila = agregar_tabla(ws, "tFranjas", ["Franja", "HoraInicio", "HoraCierre"], filas_franjas, fila) + 3

    ws.cell(row=fila, column=1, value="tTurnosComedor").font = Font(bold=True, italic=True)
    fila += 1
    filas_turnos = [(t["Turno"], t["Desde"], t["Hasta"]) for t in cfg.TURNOS_COMEDOR]
    fila = agregar_tabla(ws, "tTurnosComedor", ["Turno", "Desde", "Hasta"], filas_turnos, fila) + 3

    ws.cell(row=fila, column=1, value="tEventosValidos").font = Font(bold=True, italic=True)
    fila += 1
    filas_eventos = [(e,) for e in cfg.EVENTOS_VALIDOS]
    fila = agregar_tabla(ws, "tEventosValidos", ["Evento"], filas_eventos, fila) + 3

    # tLectores se pone más a la derecha para que quede a la vista junto con
    # las tablas chicas, sin que una tabla larga (100 filas) empuje todo lo
    # demás hacia abajo.
    filas_lectores = sorted(cfg.LECTORES.items())
    agregar_tabla(ws, "tLectores", ["Lector", "Tipo"], filas_lectores, 4, col_inicio=6)
    ws.cell(row=3, column=6, value="tLectores").font = Font(bold=True, italic=True)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["F"].width = 46
    ws.column_dimensions["G"].width = 12

    ws.sheet_view.showGridLines = False


def hoja_salida_vacia(wb, nombre, columnas, nota_extra=""):
    ws = wb.create_sheet(nombre)
    ws["A1"] = (
        f"Esta hoja se completa al cargar la consulta '{nombre}' desde Power Query "
        f"(ver GUIA_INSTALACION.md). Columnas esperadas: {', '.join(columnas)}."
    )
    ws["A1"].alignment = Alignment(wrap_text=True)
    ws["A1"].font = Font(italic=True, color="808080")
    ws.merge_cells("A1:H1")
    ws.row_dimensions[1].height = 30
    if nota_extra:
        ws["A2"] = nota_extra
        ws["A2"].font = Font(italic=True, color="808080")
        ws.merge_cells("A2:H2")
    ws.sheet_view.showGridLines = False
    return ws


def main(ruta_salida="Plantilla_Reportes_POB.xlsx"):
    wb = Workbook()
    hoja_config(wb)

    ws_dash = wb.create_sheet("Dashboard")
    ws_dash["A1"] = "Dashboard"
    ws_dash["A1"].font = Font(bold=True, size=16)
    ws_dash["A2"] = (
        "Armá acá las tablas dinámicas, segmentaciones y gráficos indicados en GUIA_INSTALACION.md "
        "(sección 6), o corré la macro Instalar_Reportes para generarlos automáticamente."
    )
    ws_dash["A2"].alignment = Alignment(wrap_text=True)
    ws_dash.merge_cells("A2:H2")
    ws_dash.sheet_view.showGridLines = False

    hoja_salida_vacia(wb, "POB_Diario",
                       ["Fecha", "Segmento", "Empresa", "Presentes_en_el_dia", "Maximo_simultaneo", "Hora_del_maximo"])
    hoja_salida_vacia(wb, "POB_Horario", ["Fecha", "Hora", "Segmento", "Empresa", "POB"])
    hoja_salida_vacia(wb, "POB_Franjas", ["Fecha", "Franja", "Segmento", "Empresa", "POB"])
    hoja_salida_vacia(wb, "Comedor", ["Fecha", "Turno", "Segmento", "Empresa", "Personas_unicas"])
    hoja_salida_vacia(wb, "Control_Calidad", ["Indicador", "Valor"])

    ws_estadias = hoja_salida_vacia(
        wb, "Estadias",
        ["IdPersona", "Segmento", "Inicio", "Fin", "Motivo", "Empresa"],
        nota_extra="Hoja de auditoría (detalle por persona). Se puede ocultar una vez cargada.",
    )
    ws_estadias.sheet_state = "hidden"

    wb.active = 0
    wb.save(ruta_salida)
    print(f"Generado: {ruta_salida}")


if __name__ == "__main__":
    main()
